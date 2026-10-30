#!/usr/bin/env python3
"""Обновляет data/prices.json для СытоПлана по открытым данным Росстата.

Источники (публичные, без авторизации):
  * https://rosstat.gov.ru/statistics/price — «Средние потребительские цены»:
      nedel_sred_cen.xlsx            — еженедельные средние цены по РФ (≈30 продуктов)
      sred_potreb_cen_MM-YYYY.xlsx   — ежемесячные средние цены по РФ и регионам (≈130 продуктов)
  TLS rosstat.gov.ru подписан Russian Trusted Root CA → используем tools/certs/ru_ca.pem
  (корневой + промежуточные сертификаты Минцифры с gu-st.ru) только для этих запросов.

Магазины (Пятёрочка, Магнит, Перекрёсток, Лента, Ашан, Дикси) не отдают каталоги анонимно
(JS-челлендж / 401 / 403 на 30.10.2026), поэтому цена магазина = цена Росстата × коэффициент
уровня цен сети (STORE_COEF ниже, источники в комментарии). Это ориентир, не цена на полке.

Запуск:  python3 tools/update_prices.py            (скачает свежие файлы)
         python3 tools/update_prices.py --local DIR (взять nedel.xlsx / mes.xlsx из DIR)
Нужно: python3, openpyxl, curl, node (чтобы прочитать каталог ингредиентов из js/data.js).
"""
import json, os, re, subprocess, sys, datetime, tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CA = os.path.join(ROOT, 'tools', 'certs', 'ru_ca.pem')
BASE = 'https://rosstat.gov.ru'

# Коэффициенты уровня цен сетей относительно средней цены (≈ средней по рынку).
# Источники: prosto.ru, сравнение корзины из 10 продуктов, Москва, апрель 2026
#   (Ашан 1276 ₽, Магнит 1411, Дикси 1517, Лента 1585, Пятёрочка 1633, Перекрёсток 1671; среднее 1515,5);
#   ngnovoros.ru, корзина из 13 продуктов, Москва, сентябрь 2025 (Пятёрочка дешевле Магнита на 12% и Ленты на 16%).
# Источники противоречат друг другу, поэтому отклонение от 1 по prosto.ru уменьшено вдвое. Значения ориентировочные.
STORE_COEF = {'p5': 1.04, 'magnit': 0.97, 'perek': 1.05, 'lenta': 1.02, 'dixy': 1.00, 'auchan': 0.92}

# ключ ингредиента → (регулярка по названию позиции Росстата, поправочный множитель).
# множитель 1.0 = прямое соответствие; ≠1 = производная оценка (например, филе дороже целой курицы).
MAP = {
 'potato':(r'^Картофель',1), 'onion':(r'^Лук репчатый',1), 'carrot':(r'^Морковь',1), 'garlic':(r'^Чеснок',1),
 'cabbage':(r'^Капуста белокочанная',1), 'beet':(r'^Свёкла|^Свекла',1), 'tomato':(r'^Помидоры',1), 'cucumber':(r'^Огурцы',1),
 'pepper':(r'^Перец сладкий',1), 'greens':(r'^Зелень',1), 'cilantro':(r'^Зелень',1.1), 'green_onion':(r'^Зелень',0.9),
 'mushrooms':(r'^Грибы',1), 'apple':(r'^Яблоки',1), 'banana':(r'^Бананы',1), 'pear':(r'^Груши',1), 'orange':(r'^Апельсины',1),
 'lemon':(r'^Лимоны',1),
 'chicken_fillet':(r'^Куры',1.6), 'chicken_thigh':(r'^Окорочка',1.1), 'chicken_drum':(r'^Окорочка',1.05), 'chicken_wings':(r'^Куры',1.0),
 'minced_chicken':(r'^Куры',1.3), 'turkey_fillet':(r'^Мясо индейки',1.1), 'minced_turkey':(r'^Мясо индейки',0.95),
 'minced_mix':(r'^Фарш мясной',1), 'beef':(r'^Говядина бескостная',1), 'pork':(r'^Свинина бескостная',1),
 'liver':(r'^Печень птицы',1), 'sausages':(r'^Сосиски',1), 'ham':(r'^Колбаса вареная',1.15),
 'fish_white':(r'^Рыба мороженая разделанная',1), 'hake':(r'^Рыба мороженая разделанная',1.15),
 'fish_red':(r'^Рыба охлажденная и мороженая разделанная лососевых',0.6), 'mackerel':(r'^Рыба мороженая неразделанная',1.1),
 'herring':(r'^Филе сельди',1), 'shrimp':(r'^Креветки',1), 'squid':(r'^Кальмары',1),
 'tuna':(r'^Консервы рыбные натуральные',1.2), 'saury':(r'^Консервы рыбные натуральные',0.9),
 'pelmeni':(r'^Пельмени',1), 'vareniki':(r'^Пельмени',0.6),
 'frozen_veg':(r'^Овощи замороженные',1), 'broccoli':(r'^Овощи замороженные',1.15), 'cauliflower':(r'^Овощи замороженные',1.05),
 'green_beans':(r'^Овощи замороженные',1), 'peas':(r'^Овощи замороженные',0.9), 'berries':(r'^Ягоды замороженные',1),
 'milk':(r'^Молоко питьевое цельное пастеризованное 2,5|^Молоко питьевое цельное пастеризованное',1), 'kefir':(r'^Кефир',1),
 'ryazhenka':(r'^Кисломолочные продукты',1), 'smetana':(r'^Сметана',1), 'tvorog':(r'^Творог,|^Творог$|^Творог',1),
 'cheese':(r'^Сыры твердые|^Сыры',1), 'mozzarella':(r'^Сыры твердые',1.1), 'processed_cheese':(r'^Сыры плавленые',1),
 'butter':(r'^Масло сливочное',1), 'yogurt':(r'^Йогурт',1), 'cream':(r'^Сливки',1), 'feta':(r'^Национальные сыры и брынза',1),
 'condensed_milk':(r'^Молоко сгущенное',1), 'eggs':(r'^Яйца',1),
 'bread':(r'^Хлеб и булочные изделия из пшеничной|^Хлеб пшеничный',1), 'bread_whole':(r'^Хлеб из ржаной|^Хлеб ржаной',1.4),
 'lavash':(r'^Хлеб и булочные изделия из пшеничной|^Хлеб пшеничный',1.4),
 'buckwheat':(r'^Крупа гречневая|^Гречка|^Крупа гречн',1), 'rice':(r'^Рис',1), 'millet':(r'^Пшено',1), 'semolina':(r'^Крупа манная',1),
 'barley':(r'^Крупы овсяная и перловая',1), 'oats':(r'^Овсяные хлопья',1), 'pasta':(r'^Макарон',1), 'spaghetti':(r'^Макарон',1.1),
 'noodles':(r'^Вермишель',2.0), 'flour':(r'^Мука пшеничная',1), 'sugar':(r'^Сахар',1), 'salt':(r'^Соль',1),
 'oil':(r'^Масло подсолнечное',1), 'spices':(r'^Сухие приправы',1),
 'black_pepper':(r'^Сухие приправы',1.6), 'bay_leaf':(r'^Сухие приправы',2.5), 'cinnamon':(r'^Сухие приправы',1.5), 'khmeli':(r'^Сухие приправы',1.0),
 'coriander':(r'^Сухие приправы',1.0), 'paprika':(r'^Сухие приправы',1.2), 'cumin':(r'^Сухие приправы',1.6), 'curry':(r'^Сухие приправы',1.2),
 'turkey_thigh':(r'^Мясо индейки',0.85), 'minced_beef':(r'^Фарш мясной',1.15), 'pork_neck':(r'^Свинина бескостная',1.1), 'minced_pork':(r'^Фарш мясной',0.95),
 'pork_ribs':(r'^Свинина \(кроме бескостного',1.0), 'lasagna':(r'^Макарон',4.5), 'tomato_paste':(r'^Консервы томатные',1),
 'canned_tomatoes':(r'^Консервы томатные',0.8), 'ketchup':(r'^Кетчуп',1), 'mayo':(r'^Майонез',1), 'honey':(r'^Мёд|^Мед',1),
 'jam':(r'^Варенье',1), 'nuts':(r'^Орехи',1), 'raisins':(r'^Сухофрукты',0.7), 'prunes':(r'^Сухофрукты',1), 'dried_apricots':(r'^Сухофрукты',1.1),
 'lentils':(r'^Горох и фасоль',1.8), 'split_peas':(r'^Горох и фасоль',0.8), 'chickpeas':(r'^Горох и фасоль',1.9),
 'canned_beans':(r'^Овощи натуральные консервированные',1), 'corn':(r'^Овощи натуральные консервированные',1),
 'green_peas':(r'^Горошек зеленый',1), 'pickles':(r'^Овощи натуральные консервированные',0.9),
}
PIECE_G = {'lemon': 130}  # средний вес одной штуки для перевода цены за кг в цену за шт
MONTHS = {'января':1,'февраля':2,'марта':3,'апреля':4,'мая':5,'июня':6,'июля':7,'августа':8,'сентября':9,'октября':10,'ноября':11,'декабря':12}

def curl(url, out=None):
    cmd = ['curl', '-sSfL', '--max-time', '180', '--cacert', CA, '-A', 'SytoPlan-price-updater/1.0', url]
    if out: cmd += ['-o', out]
    r = subprocess.run(cmd, capture_output=True, check=True)
    return r.stdout.decode('utf-8', 'replace') if not out else out

def catalog():
    js = ("const fs=require('fs'),vm=require('vm');const c={console};vm.createContext(c);"
          "vm.runInContext(fs.readFileSync(process.argv[1],'utf8').replace(/^const /gm,'var '),c);"
          "console.log(JSON.stringify(Object.fromEntries(Object.entries(c.ING).map(([k,g])=>[k,{name:g.name,pack:g.pack,unit:g.unit,price:g.price,w:g.w}]))))")
    return json.loads(subprocess.run(['node', '-e', js, os.path.join(ROOT, 'js', 'data.js')], capture_output=True, check=True).stdout)

def weekly(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sorted([s for s in wb.sheetnames if s.isdigit()])[-1]] if any(s.isdigit() for s in wb.sheetnames) else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True)); hdr = None; out = {}
    for r in rows:
        if r and r[0] and str(r[0]).strip().startswith('Наименование'): hdr = r; continue
        if hdr and r and r[0]:
            vals = [(i, v) for i, v in enumerate(r) if i > 0 and isinstance(v, (int, float))]
            if vals: i, v = vals[-1]; out[str(r[0]).strip()] = (float(v), str(hdr[i]).strip())
    year = ws.title if ws.title.isdigit() else str(datetime.date.today().year)
    return out, year

def monthly(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[[s for s in wb.sheetnames if re.match(r'\d\d\(\d{4}\)', s)][-1]]
    names = None; out = {}
    for r in ws.iter_rows(values_only=True):
        if names is None and r[0] is None and r[2] and isinstance(r[2], str) and ',' in r[2]: names = r; continue
        if names and r[0] is not None and str(r[0]).strip() == '643':
            for i, v in enumerate(r):
                if i >= 2 and isinstance(v, (int, float)) and names[i]: out[str(names[i]).strip()] = float(v)
            break
    m = re.match(r'(\d\d)\((\d{4})\)', ws.title)
    return out, f'{m.group(2)}-{m.group(1)}'

def find(table, rx):
    for name in table:
        if re.search(rx, name): return name
    return None

def main():
    tmp = tempfile.mkdtemp()
    if '--local' in sys.argv:
        d = sys.argv[sys.argv.index('--local') + 1]; wf, mf = os.path.join(d, 'nedel.xlsx'), os.path.join(d, 'mes.xlsx')
    else:
        page = curl(BASE + '/statistics/price')
        wl = re.search(r'href="(/storage/mediabank/nedel_sred_cen[^"]*\.xlsx)"', page).group(1)
        ml = re.search(r'href="(/storage/mediabank/sred_potreb_cen_\d\d-\d{4}\.xlsx)"', page).group(1)
        wf = curl(BASE + wl, os.path.join(tmp, 'nedel.xlsx')); mf = curl(BASE + ml, os.path.join(tmp, 'mes.xlsx'))
    W, wyear = weekly(wf); M, mper = monthly(mf)
    wdate_raw = next(iter(W.values()))[1] if W else ''
    mm = re.search(r'(\d+)\s+(\w+)', wdate_raw)
    wdate = f'{wyear}-{MONTHS.get(mm.group(2), 1):02d}-{int(mm.group(1)):02d}' if mm else wyear
    cat = catalog(); res = {}; stats = {'rosstat': 0, 'derived': 0, 'demo': 0}
    for k, g in cat.items():
        rec = None
        if k in MAP:
            rx, f = MAP[k]
            wn = find(W, rx); mn = find(M, rx)
            if wn: per, src, item, date = W[wn][0], 'rosstat-weekly', wn, wdate
            elif mn: per, src, item, date = M[mn], 'rosstat-monthly', mn, mper
            else: per = None
            if per:
                if g['unit'] == 'шт' and k == 'eggs': base = per * g['pack'] / 10
                elif g['unit'] == 'шт': base = per * PIECE_G.get(k, g['w'] or 100) / 1000 * g['pack']
                else: base = per * g['pack'] / 1000
                base *= f
                rec = {'base': round(base, 2), 'source': src if f == 1 else src + '-derived', 'item': item,
                       'per_unit': round(per, 2), 'factor': f, 'date': date}
                stats['rosstat' if f == 1 else 'derived'] += 1
        if not rec:
            rec = {'base': g['price'], 'source': 'demo', 'item': None, 'date': None}; stats['demo'] += 1
        for s, c in STORE_COEF.items(): rec[s] = int(round(rec['base'] * c))
        res[k] = rec
    n = len(cat)
    res['_meta'] = {
        'updated': datetime.date.today().isoformat(), 'weekly_date': wdate, 'monthly_period': mper,
        'region': 'Российская Федерация (средние цены)', 'store_coef': STORE_COEF,
        'coverage': {'ingredients': n, 'rosstat_direct': stats['rosstat'], 'rosstat_derived': stats['derived'], 'demo': stats['demo'],
                     'real_pct': round(100 * (stats['rosstat'] + stats['derived']) / n, 1), 'store_direct_pct': 0},
        'sources': ['https://rosstat.gov.ru/statistics/price (nedel_sred_cen.xlsx, sred_potreb_cen)',
                    'store_coef: prosto.ru (апрель 2026), ngnovoros.ru (сентябрь 2025) — ориентировочно'],
        'note': 'Цена магазина = средняя цена Росстата × коэффициент сети; не цена на полке.'}
    os.makedirs(os.path.join(ROOT, 'data'), exist_ok=True)
    with open(os.path.join(ROOT, 'data', 'prices.json'), 'w', encoding='utf-8') as fh:
        json.dump(res, fh, ensure_ascii=False, indent=1)
    print(json.dumps(res['_meta']['coverage'], ensure_ascii=False), 'weekly', wdate, 'monthly', mper)

if __name__ == '__main__':
    main()
